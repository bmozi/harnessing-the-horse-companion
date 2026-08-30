#!/usr/bin/env python3
# © 2026 John Briggs - MIT licensed
"""Build an accessible DOCX and tagged PDF edition of the five teaching panels."""

from __future__ import annotations

import shutil
import sys
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
VERSION = json.loads((ROOT / "package.json").read_text(encoding="utf-8"))["version"]
OUTPUT_DOCX = ROOT / "output" / "docx" / f"Harnessing-the-Horse-Software-Factory-Teaching-Panels-v{VERSION}.docx"
OUTPUT_PDF_DIR = ROOT / "output" / "pdf"

NAVY = "11263D"
BLUE = "1768AC"
MUTED = "607384"

PANELS = [
    {
        "title": "Evidence before build",
        "lead": "A work order earns an agent turn only after scope, evidence, safety, workspace, and cost boundaries are visible.",
        "description": "The flow moves from a bounded work order through a deterministic evidence-and-safety preflight to one of three explicit outcomes: proceed, narrow the work, or request human review. The work order remains the accountability boundary in both directions.",
        "image": "diagrams/merlin-architecture-v2-visual-1.png",
    },
    {
        "title": "One agent. Six disciplined phases.",
        "lead": "The EXPRESS delivery arc preserves intent by keeping one accountable agent responsible from understanding through self-audit.",
        "description": "The six phases are Understand, Investigate, Plan, Test, Implement, and Self-audit. Continuity protects design intent and local context, but it does not replace deterministic checks, independent verification, cost boundaries, or human judgment.",
        "image": "diagrams/merlin-architecture-v2-visual-2.png",
    },
    {
        "title": "Proof is not one undifferentiated gate",
        "lead": "Verification is classified by consequence so rigor increases trust without turning every signal into critical-path ceremony.",
        "description": "Blocking checks stop unsafe work; conditional checks become blocking when relevant; advisory checks inform judgment; asynchronous checks improve later work. Independent challenge and bounded repair keep the implementer from being the only judge of its own claim.",
        "image": "diagrams/merlin-architecture-v2-visual-3.png",
    },
    {
        "title": "Automation earns trust by showing its work",
        "lead": "State, evidence, cost, failures, and next actions remain visible so a person can intervene without reconstructing the run.",
        "description": "The agent produces evidence, the control plane keeps work convergent, and the operator decides with context. Durable state, observability, cost safety, and workspace isolation make human review a healthy architectural outcome rather than a hidden failure.",
        "image": "diagrams/merlin-architecture-v2-visual-4.png",
    },
    {
        "title": "Every run can improve the next",
        "lead": "Learning compounds only when it is verified, traceable, measured, and kept outside the delivery path whenever possible.",
        "description": "Outcome signals become verified memory, which can propose a bounded improvement. Each candidate must be evaluated and either promoted, discarded, or rolled back. The reusable rules are to bound work, preserve accountability, classify proof, escalate judgment, expose state and cost, and promote learning only after evidence.",
        "image": "diagrams/merlin-architecture-v2-visual-5.png",
    },
]


def set_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def build_docx() -> Path:
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.34)
    section.bottom_margin = Inches(0.34)
    section.left_margin = Inches(0.42)
    section.right_margin = Inches(0.42)
    section.header_distance = Inches(0.16)
    section.footer_distance = Inches(0.16)

    normal = document.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(8.7)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0
    heading = document.styles["Heading 1"]
    heading.font.name = "Liberation Sans"
    heading.font.size = Pt(19)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(NAVY)
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True

    props = document.core_properties
    props.title = "Harnessing the Horse - Software Factory Teaching Panels"
    props.subject = "Five accessible teaching panels for the public Merlin Software Factory architecture."
    props.author = "John Briggs"
    props.keywords = "AI engineering, software factory, EXPRESS, verification, human judgment, accessible diagrams"
    props.comments = f"Accessible reader edition, companion release v{VERSION}."

    header = section.header.paragraphs[0]
    header.text = "HARNESSING THE HORSE  |  SOFTWARE FACTORY TEACHING PANELS"
    header.runs[0].font.size = Pt(7)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.add_run("Public teaching architecture  |  Panel ")
    page_number(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in footer.runs:
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    for index, panel in enumerate(PANELS):
        title = document.add_heading(panel["title"], level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lead = document.add_paragraph(panel["lead"])
        lead.runs[0].bold = True
        lead.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
        lead.paragraph_format.space_after = Pt(2)
        description = document.add_paragraph(panel["description"])
        description.runs[0].font.size = Pt(8.2)
        description.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
        description.paragraph_format.space_after = Pt(3)
        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_after = Pt(0)
        shape = image_paragraph.add_run().add_picture(str(ROOT / panel["image"]), width=Inches(6.55))
        set_alt_text(shape, f"Panel {index + 1}: {panel['title']}", f"{panel['lead']} {panel['description']}")
        if index < len(PANELS) - 1:
            document.add_page_break()

    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


def main() -> int:
    docx_path = build_docx()
    sys.path.insert(0, str(ROOT / "scripts"))
    from build_reader_guide import convert_to_pdf

    pdf_path = convert_to_pdf(docx_path, OUTPUT_PDF_DIR)
    shutil.copy2(pdf_path, ROOT / "diagrams" / "merlin-architecture-v2-visual.pdf")
    print(docx_path)
    print(pdf_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
