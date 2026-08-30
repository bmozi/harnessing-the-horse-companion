#!/usr/bin/env python3
# © 2026 John Briggs - MIT licensed
"""Build the semantic DOCX and tagged PDF Reader Quick Start."""

from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "release-assets" / "reader-quick-start.json"
DOCX_DIR = ROOT / "output" / "docx"
PDF_DIR = ROOT / "output" / "pdf"

NAVY = "11263D"
TEAL = "168C8C"
INK = "243447"
MUTED = "607384"
PALE = "EAF4F3"
LIGHT = "F3F6F8"
WHITE = "FFFFFF"


def set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def shade_paragraph(paragraph: Any, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_cell_margins(cell: Any, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_alt_text(inline_shape: Any, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def page_number(paragraph: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def add_hyperlink(paragraph: Any, label: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([run_properties, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_document(document: Document, title: str, subject: str, version: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(9.8)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.06

    for name, size, color in (("Title", 31, NAVY), ("Heading 1", 22, NAVY), ("Heading 2", 12, TEAL)):
        style = styles[name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(2)
        style.paragraph_format.space_after = Pt(7)

    styles["Subtitle"].font.name = "Liberation Sans"
    styles["Subtitle"].font.size = Pt(17)
    styles["Subtitle"].font.color.rgb = RGBColor.from_string(TEAL)

    props = document.core_properties
    props.title = title
    props.subject = subject
    props.author = "John Briggs"
    props.keywords = "AI engineering, agentic development, accountable software delivery, reader guide"
    props.comments = f"Accessible reader edition, companion release v{version}."

    header = section.header.paragraphs[0]
    header.text = "HARNESSING THE HORSE  |  READER QUICK START"
    header.runs[0].font.size = Pt(7.5)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer.paragraphs[0]
    footer.add_run(f"Companion v{version}   |   ")
    page_number(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in footer.runs:
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def add_callout(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    shade_paragraph(paragraph, "FFF3E8")
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    run.font.size = Pt(9.6)


def add_step_table(document: Document, rows: list[list[str]]) -> None:
    for index, (label, detail) in enumerate(rows):
        paragraph = document.add_paragraph()
        shade_paragraph(paragraph, PALE if index % 2 == 0 else LIGHT)
        paragraph.paragraph_format.left_indent = Inches(0.1)
        paragraph.paragraph_format.right_indent = Inches(0.1)
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}: ")
        label_run.bold = True
        label_run.font.color.rgb = RGBColor.from_string(TEAL)
        paragraph.add_run(detail)


def add_artifact_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(1.26), Inches(1.68), Inches(4.24))
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell, 75, 80, 75, 80)
            shade(cell, NAVY if row_index == 0 else (WHITE if row_index % 2 else LIGHT))
            run = cell.paragraphs[0].add_run(value)
            run.font.size = Pt(8.3 if row_index else 8.5)
            run.bold = row_index == 0
            run.font.color.rgb = RGBColor.from_string(WHITE if row_index == 0 else INK)
    set_repeat_table_header(table.rows[0])


def add_page(document: Document, page: dict[str, Any]) -> None:
    document.add_heading(page["title"], level=1)
    kicker = document.add_paragraph(page["kicker"])
    kicker.paragraph_format.space_after = Pt(8)
    kicker.runs[0].italic = True
    kicker.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    if "image" in page:
        image_path = ROOT / page["image"]
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = paragraph.add_run().add_picture(str(image_path), width=Inches(5.5))
        set_alt_text(
            shape,
            "Governed AI software factory overview",
            "A layered software delivery system in which bounded intake enters an accountable delivery arc, independent evidence supports human decisions, and verified learning improves later work.",
        )
    if "table" in page:
        add_artifact_table(document, page["table"])
    if "steps" in page:
        add_step_table(document, page["steps"])
    if "bullets" in page:
        for item in page["bullets"]:
            paragraph = document.add_paragraph(item, style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(1.5)
    add_callout(document, page["callout"])


def convert_to_pdf(docx_path: Path, pdf_dir: Path) -> Path:
    converter = shutil.which("soffice") or shutil.which("libreoffice")
    if not converter:
        raise RuntimeError("LibreOffice is required to create the tagged PDF")
    pdf_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="horse-docx-profile-") as profile:
        subprocess.run(
            [converter, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(pdf_dir), str(docx_path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
    pdf_path = pdf_dir / f"{docx_path.stem}.pdf"
    if not pdf_path.exists():
        raise RuntimeError(f"LibreOffice did not produce {pdf_path}")
    return pdf_path


def build() -> tuple[Path, Path]:
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    version = data["version"]
    basename = f"Harnessing-the-Horse-Reader-Quick-Start-v{version}"
    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    docx_path = DOCX_DIR / f"{basename}.docx"

    document = Document()
    configure_document(
        document,
        f"{data['title']} - {data['subtitle']}",
        "A practical orientation to the Harnessing the Horse companion repository.",
        version,
    )

    title = document.add_heading(data["title"], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(data["subtitle"], style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    promise = document.add_paragraph(data["promise"])
    promise.alignment = WD_ALIGN_PARAGRAPH.CENTER
    promise.paragraph_format.left_indent = Inches(0.68)
    promise.paragraph_format.right_indent = Inches(0.68)
    promise.paragraph_format.space_before = Pt(72)
    promise.paragraph_format.space_after = Pt(72)
    promise.runs[0].font.size = Pt(14)
    promise.runs[0].font.bold = True
    promise.runs[0].font.color.rgb = RGBColor.from_string(NAVY)
    byline = document.add_paragraph(f"John Briggs\nCompanion release v{version} | August 29, 2026")
    byline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    byline.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    document.add_page_break()

    for index, page in enumerate(data["pages"]):
        add_page(document, page)
        if index < len(data["pages"]) - 1:
            document.add_page_break()

    document.add_heading("Continue with the stable resources", level=2)
    for label, url in data["links"]:
        paragraph = document.add_paragraph()
        add_hyperlink(paragraph, label, url)
        trailing = paragraph.add_run(f" - {url}")
        trailing.font.size = Pt(8)
        trailing.font.color.rgb = RGBColor.from_string(MUTED)
    license_paragraph = document.add_paragraph("Written content © 2026 John Briggs. Licensed under CC BY-NC-SA 4.0.")
    license_paragraph.runs[0].font.size = Pt(8)
    license_paragraph.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    document.save(docx_path)
    pdf_path = convert_to_pdf(docx_path, PDF_DIR)
    return docx_path, pdf_path


if __name__ == "__main__":
    for built in build():
        print(built)
